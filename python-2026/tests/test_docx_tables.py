"""Regression coverage for DOCX table flattening and structured ingestion."""

from io import BytesIO

from docx import Document
from services.ingestion.items import expand_section_to_items
from services.ingestion.tagger import tag_role
from utils.chunker import chunk_spec_recursive
from utils.docx_reader import extract_text_from_docx, flatten_table_records


def _document_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_horizontal_table_rows_become_atomic_heading_aware_items() -> None:
    document = Document()
    document.add_heading("Product specification", level=1)
    document.add_paragraph("General product introduction.")
    document.add_heading("Actors", level=2)

    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Actor"
    table.cell(0, 1).text = "Actions"
    table.cell(1, 0).text = "Admin"
    table.cell(1, 1).text = "Create users\nDeactivate accounts"
    table.cell(2, 0).text = "User"
    table.cell(2, 1).text = "Update profile"
    document.add_paragraph("Additional actor guidance remains in this section.")

    chunks = chunk_spec_recursive(document)
    actor_chunk = next(
        chunk
        for chunk in chunks
        if chunk.heading_path == ["Product specification", "Actors"]
    )

    assert actor_chunk.text.splitlines() == [
        "- Actor: Admin | Actions: Create users; Deactivate accounts",
        "- Actor: User | Actions: Update profile",
        "Additional actor guidance remains in this section.",
    ]

    items = expand_section_to_items(actor_chunk)
    assert [item.text for item in items[:2]] == [
        "Actor: Admin | Actions: Create users; Deactivate accounts",
        "Actor: User | Actions: Update profile",
    ]
    tag_role(items)
    assert [item.role for item in items[:2]] == ["ACTOR", "ACTOR"]


def test_vertical_key_value_table_becomes_one_record() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Actor"
    table.cell(0, 1).text = "Admin"
    table.cell(1, 0).text = "Actions"
    table.cell(1, 1).text = "Create users and deactivate accounts"

    assert flatten_table_records(table) == [
        "Actor: Admin | Actions: Create users and deactivate accounts"
    ]


def test_plain_text_extraction_uses_the_same_labelled_rows() -> None:
    document = Document()
    document.add_heading("Actors", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Actor"
    table.cell(0, 1).text = "Actions"
    table.cell(1, 0).text = "Admin"
    table.cell(1, 1).text = "Create users"

    text = extract_text_from_docx(_document_bytes(document))

    assert "# Actors" in text
    assert "- Actor: Admin | Actions: Create users" in text
    assert "Actor | Actions" not in text
